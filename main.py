"""
NESHER - Liberação de Chapas
Backend que liga a tela web à planilha real (Controle_Chapas_MDF_FUNDOS.xlsx),
lendo e gravando na aba "CHAPAS POR LOTE" (Tabela10).

Como rodar:
    pip install -r requirements.txt
    python main.py
Depois abra http://localhost:8000 no navegador.

Por padrão ele usa o arquivo Controle_Chapas_MDF_FUNDOS.xlsx que está nesta mesma
pasta. Se a planilha real estiver em outro lugar (ex: um servidor da empresa),
aponte para ela com a variável de ambiente PLANILHA_PATH, por exemplo:

    PLANILHA_PATH="//SERVIDOR/PPCP/Controle_Chapas_MDF_FUNDOS.xlsx" python main.py
"""
import os
from datetime import datetime, date
from typing import Optional

from fastapi import FastAPI, HTTPException
from fastapi.middleware.cors import CORSMiddleware
from fastapi.staticfiles import StaticFiles
from fastapi.responses import FileResponse
from pydantic import BaseModel
import openpyxl
from openpyxl.utils import range_boundaries
from copy import copy

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
PLANILHA_PATH = os.environ.get(
    "PLANILHA_PATH", os.path.join(BASE_DIR, "Controle_Chapas_MDF_FUNDOS.xlsx")
)
SHEET_NAME = "CHAPAS POR LOTE"
TABLE_NAME = "Tabela10"

# Espelham as abas de referência da própria planilha (REFERÊNCIA / LISTA),
# usadas para calcular Fábrica e M² da chapa sem depender de fórmulas do Excel.
PREFIXO_FABRICA = {
    "N": "Nesher",
    "I": "Multilínea",
    "M": "Multilínea",
    "L": "Líder",
    "U": "Líder",
    "A": "Nesher",
    "E": "Nesher",
    "V": "Via",
}

M2_POR_TIPO_CHAPA = {
    "MDF - 9MM": 5.0875,
    "MDF - 15MM": 6.05,
    "MDF - 25MM 2750x2200": 6.05,
    "MDF - 25MM 2750x2100": 5.775,
    "EUCADUR BR - 2,5MM": 5.8575,
    "EUCADUR FRJ - 2,5MM": 5.8575,
    "EUCADUR NT - 2,5MM": 5.8575,
    "CAPA DURATREE -2,5MM": 5.8575,
    "CAPA MDF - 15MM": 6.05,
    "MDF - 18MM 2750X1850": 5.0875,
    "CAPA MDF - 12MM": 6.05,
    "MDF - 25MM 2750X450 - SOBRAS": 1.2375,
    "TIRA 9MM - 1850X165X9": 0.30525,
    "MDF BERNECK - 15MM ": 6.05,
}

app = FastAPI(title="NESHER - Liberação de Chapas")
app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_methods=["*"],
    allow_headers=["*"],
)


class NovaLiberacao(BaseModel):
    op: str
    tipo_chapa: str
    qtde_chapas: float
    perda_pct: Optional[float] = None
    erp: Optional[float] = None


def fabrica_por_op(op: str) -> str:
    prefixo = (op or "").strip()[:1].upper()
    return PREFIXO_FABRICA.get(prefixo, "—")


def _valor_data(v):
    if hasattr(v, "isoformat"):
        return v.isoformat()
    return v


def carregar_linhas():
    """Lê a tabela Tabela10 (aba CHAPAS POR LOTE) direto do arquivo atual."""
    if not os.path.exists(PLANILHA_PATH):
        raise HTTPException(500, f"Planilha não encontrada em: {PLANILHA_PATH}")
    wb = openpyxl.load_workbook(PLANILHA_PATH, data_only=True)
    ws = wb[SHEET_NAME]
    tbl = ws.tables[TABLE_NAME]
    min_col, min_row, max_col, max_row = range_boundaries(tbl.ref)
    header_row = min_row

    linhas = []
    for r in range(header_row + 1, max_row + 1):
        vals = [ws.cell(row=r, column=c).value for c in range(min_col, max_col + 1)]
        if all(v is None for v in vals):
            continue
        vals = (vals + [None] * 12)[:12]
        (
            data, op, fabrica, tipo, m2_chapa, qtde, perda,
            m2total, erp, chapas_erp, dif_m2, dif_chapas,
        ) = vals
        linhas.append(
            {
                "data": _valor_data(data),
                "op": op,
                "fabrica": fabrica,
                "tipo_chapa": tipo,
                "m2_chapa": m2_chapa,
                "qtde_chapas_liberadas": qtde,
                "perda_pct": perda,
                "m2_total": m2total,
            }
        )
    return linhas


@app.get("/api/liberacoes")
def listar_liberacoes(limit: int = 50):
    linhas = carregar_linhas()
    linhas.sort(key=lambda x: x["data"] or "", reverse=True)
    return linhas[:limit]


@app.get("/api/stats")
def stats():
    linhas = carregar_linhas()
    hoje = date.today().isoformat()
    hoje_linhas = [l for l in linhas if (l["data"] or "").startswith(hoje)]
    return {
        "chapas_liberadas_hoje": sum(l["qtde_chapas_liberadas"] or 0 for l in hoje_linhas),
        "ops_atendidas_hoje": len({l["op"] for l in hoje_linhas if l["op"]}),
        "total_chapas_liberadas": sum(l["qtde_chapas_liberadas"] or 0 for l in linhas),
        "total_registros": len(linhas),
    }


@app.get("/api/tipos-chapa")
def tipos_chapa():
    return sorted(M2_POR_TIPO_CHAPA.keys())


@app.post("/api/liberacoes")
def criar_liberacao(nova: NovaLiberacao):
    if nova.tipo_chapa not in M2_POR_TIPO_CHAPA:
        raise HTTPException(400, f"Tipo de chapa desconhecido: {nova.tipo_chapa}")
    if not nova.op or not nova.op.strip():
        raise HTTPException(400, "OP é obrigatória")
    if nova.qtde_chapas is None or nova.qtde_chapas <= 0:
        raise HTTPException(400, "Quantidade de chapas liberadas deve ser maior que zero")

    wb = openpyxl.load_workbook(PLANILHA_PATH)  # sem data_only: preserva fórmulas das outras linhas
    ws = wb[SHEET_NAME]
    tbl = ws.tables[TABLE_NAME]
    min_col, min_row, max_col, max_row = range_boundaries(tbl.ref)
    nova_linha = max_row + 1

    m2_chapa = M2_POR_TIPO_CHAPA[nova.tipo_chapa]
    fabrica = fabrica_por_op(nova.op)
    m2_total = round(m2_chapa * nova.qtde_chapas, 6)
    erp = nova.erp
    chapas_erp = round(erp / m2_chapa, 6) if erp else None
    dif_m2 = round(m2_total - erp, 6) if erp is not None else None
    dif_chapas = round(dif_m2 / m2_chapa, 6) if dif_m2 is not None else None

    valores = [
        datetime.now(), nova.op.strip(), fabrica, nova.tipo_chapa, m2_chapa,
        nova.qtde_chapas, nova.perda_pct, m2_total, erp, chapas_erp, dif_m2, dif_chapas,
    ]

    for offset, valor in enumerate(valores):
        col = min_col + offset
        cell = ws.cell(row=nova_linha, column=col, value=valor)
        # copia o formato de número da linha anterior, pra manter o padrão da planilha
        origem = ws.cell(row=max_row, column=col)
        cell.number_format = origem.number_format

    # estende a Tabela do Excel pra incluir a nova linha
    novo_ref = f"{tbl.ref.split(':')[0]}:{ws.cell(row=nova_linha, column=max_col).coordinate}"
    tbl.ref = novo_ref

    wb.save(PLANILHA_PATH)

    return {
        "ok": True,
        "linha": {
            "data": valores[0].isoformat(),
            "op": valores[1],
            "fabrica": valores[2],
            "tipo_chapa": valores[3],
            "m2_chapa": valores[4],
            "qtde_chapas_liberadas": valores[5],
            "perda_pct": valores[6],
            "m2_total": valores[7],
        },
    }


# serve a tela (index.html) na raiz do site
app.mount("/", StaticFiles(directory=os.path.join(BASE_DIR, "static"), html=True), name="static")


if __name__ == "__main__":
    import uvicorn

    uvicorn.run(app, host="0.0.0.0", port=8000)
